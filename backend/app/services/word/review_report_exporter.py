"""
审查报告 Word 文档导出器

生成结构化的 Word 格式审查报告，包含：
- 审查概述（合同名称、审查时间、置信度、风险等级）
- 风险条款列表
- 缺失条款列表
- 修改建议
- 规则判定结果
"""
import io
from datetime import datetime
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import RGBColor, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class ReviewReportExporter:
    """审查报告 Word 导出器"""

    # 风险等级颜色
    RISK_COLORS = {
        "high": RGBColor(0xC6, 0x28, 0x28),   # 红色
        "medium": RGBColor(0xE6, 0x51, 0x00), # 橙色
        "low": RGBColor(0x15, 0x65, 0xC0),    # 蓝色
    }

    def __init__(self):
        pass

    def generate_review_report(
        self,
        contract_name: str,
        review_result: Dict[str, Any],
        options: Dict[str, Any]
    ) -> bytes:
        """
        生成审查报告 Word 文档

        Args:
            contract_name: 合同文件名
            review_result: 审查结果字典
            options: 导出选项

        Returns:
            Word 文档字节流
        """
        doc = Document()

        # 设置文档样式
        self._setup_document_style(doc)

        # 添加标题
        self._add_title(doc, contract_name)

        # 添加审查概述
        self._add_overview(doc, review_result)

        # 添加风险条款
        if options.get('include_risk_clauses', True):
            self._add_risk_clauses(doc, review_result)

        # 添加缺失条款
        if options.get('include_missing_clauses', True):
            self._add_missing_clauses(doc, review_result)

        # 添加修改建议
        if options.get('include_suggestions', True):
            self._add_suggestions(doc, review_result)

        # 添加规则判定结果
        if options.get('include_rule_judgments', True):
            self._add_rule_judgments(doc, review_result)

        # 添加政策参考
        if options.get('include_policy_references', True):
            self._add_policy_references(doc, review_result)

        # 保存文档
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output.read()

    def _setup_document_style(self, doc: Document) -> None:
        """设置文档样式"""
        # 设置默认字体
        style = doc.styles['Normal']
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    def _add_title(self, doc: Document, contract_name: str) -> None:
        """添加报告标题"""
        # 主标题
        title = doc.add_heading('合同审查报告', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.color.rgb = RGBColor(0x21, 0x21, 0x21)

        # 副标题 - 合同名称
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f'合同文件：{contract_name}')
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x42, 0x42, 0x42)

        # 审查时间
        time_para = doc.add_paragraph()
        time_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        time_run = time_para.add_run(f'审查时间：{datetime.now().strftime("%Y年%m月%d日 %H:%M")}')
        time_run.font.size = Pt(10)
        time_run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

        doc.add_paragraph()  # 空行

    def _add_overview(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加审查概述"""
        # 小标题
        heading = doc.add_heading('一、审查概述', level=1)

        # 创建概述表格
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'

        # 填充数据
        data = [
            ('整体风险等级', self._get_risk_level_label(review_result.get('overall_risk_level', 'low'))),
            ('置信度评分', f"{round(review_result.get('confidence_score', 0) * 100)}%"),
            ('风险条款数量', f"{len(review_result.get('risk_clauses', []))} 条"),
            ('缺失条款数量', f"{len(review_result.get('missing_clauses', []))} 条"),
        ]

        for i, (label, value) in enumerate(data):
            row = table.rows[i]
            row.cells[0].text = label
            row.cells[1].text = value

            # 设置标签列样式
            for para in row.cells[0].paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x42, 0x42, 0x42)

        doc.add_paragraph()  # 空行

        # 审批流程说明
        if review_result.get('approval_flow_description'):
            flow_para = doc.add_paragraph()
            flow_run = flow_para.add_run('审批流程：')
            flow_run.font.bold = True
            flow_para.add_run(review_result.get('approval_flow_description'))

    def _add_risk_clauses(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加风险条款部分"""
        heading = doc.add_heading('二、风险条款', level=1)

        risk_clauses = review_result.get('risk_clauses', [])
        if not risk_clauses:
            doc.add_paragraph('无')
            return

        for i, clause in enumerate(risk_clauses, 1):
            self._add_risk_clause_item(doc, i, clause)

    def _add_risk_clause_item(self, doc: Document, index: int, clause: Dict[str, Any]) -> None:
        """添加单个风险条款"""
        # 条款标题
        title = doc.add_paragraph()
        title_run = title.add_run(f'{index}. {clause.get("title", f"风险条款 {index}")}')
        title_run.font.bold = True
        title_run.font.size = Pt(12)

        # 风险等级标签
        risk_level = clause.get('risk_level', 'medium')
        level_run = title.add_run(f'  [{self._get_risk_level_label(risk_level)}]')
        level_run.font.color.rgb = self.RISK_COLORS.get(risk_level, self.RISK_COLORS['medium'])
        level_run.font.size = Pt(10)

        # 原文本
        if clause.get('original_text'):
            p = doc.add_paragraph()
            p.style = 'Quote'
            run = p.add_run(f'原文：{clause.get("original_text")}')
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

        # 风险描述
        if clause.get('risk_description'):
            p = doc.add_paragraph()
            p.add_run('风险描述：').bold = True
            p.add_run(clause.get('risk_description'))

        # 修改建议
        if clause.get('suggestion'):
            p = doc.add_paragraph()
            p.add_run('修改建议：').bold = True
            p.add_run(clause.get('suggestion'))

        # 法律依据
        if clause.get('legal_reference'):
            p = doc.add_paragraph()
            p.add_run('法律依据：').bold = True
            run = p.add_run(clause.get('legal_reference'))
            run.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)

        # 政策参考
        if clause.get('policy_reference'):
            p = doc.add_paragraph()
            p.add_run('政策参考：').bold = True
            run = p.add_run(clause.get('policy_reference'))
            run.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)

        doc.add_paragraph()  # 空行

    def _add_missing_clauses(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加缺失条款部分"""
        heading = doc.add_heading('三、缺失条款', level=1)

        missing_clauses = review_result.get('missing_clauses', [])
        if not missing_clauses:
            doc.add_paragraph('无')
            return

        for i, clause in enumerate(missing_clauses, 1):
            self._add_missing_clause_item(doc, i, clause)

    def _add_missing_clause_item(self, doc: Document, index: int, clause: Dict[str, Any]) -> None:
        """添加单个缺失条款"""
        # 条款标题
        title = doc.add_paragraph()
        title_run = title.add_run(f'{index}. {clause.get("title", f"缺失条款 {index}")}')
        title_run.font.bold = True
        title_run.font.size = Pt(12)
        title_run.font.color.rgb = RGBColor(0xE6, 0x51, 0x00)  # 橙色

        # 描述
        if clause.get('description'):
            p = doc.add_paragraph()
            p.add_run('说明：').bold = True
            p.add_run(clause.get('description'))

        # 建议
        if clause.get('suggestion'):
            p = doc.add_paragraph()
            p.add_run('建议：').bold = True
            p.add_run(clause.get('suggestion'))

        # 法律依据
        if clause.get('legal_reference'):
            p = doc.add_paragraph()
            p.add_run('法律依据：').bold = True
            run = p.add_run(clause.get('legal_reference'))
            run.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)

        doc.add_paragraph()  # 空行

    def _add_suggestions(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加修改建议部分"""
        heading = doc.add_heading('四、修改建议', level=1)

        suggestions = review_result.get('suggestions', [])
        if not suggestions:
            doc.add_paragraph('无')
            return

        for i, suggestion in enumerate(suggestions, 1):
            p = doc.add_paragraph()
            p.add_run(f'{i}. ').bold = True
            p.add_run(suggestion.get('title', f'建议 {i}'))

            if suggestion.get('content'):
                p = doc.add_paragraph()
                p.style = 'List Bullet'
                p.add_run(suggestion.get('content'))

        doc.add_paragraph()

    def _add_rule_judgments(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加规则判定结果"""
        heading = doc.add_heading('五、规则判定结果', level=1)

        rule_judgments = review_result.get('rule_judgments', {})
        if not rule_judgments:
            doc.add_paragraph('无')
            return

        # 必须修改（违反政策）
        violations = rule_judgments.get('violations', [])
        if violations:
            sub_heading = doc.add_heading('5.1 违反政策 - 必须修改', level=2)
            for i, item in enumerate(violations, 1):
                self._add_rule_judgment_item(doc, i, item, 'high')

        # 需要审批
        approvals_needed = rule_judgments.get('approvals_needed', [])
        if approvals_needed:
            sub_heading = doc.add_heading('5.2 需要审批', level=2)
            for i, item in enumerate(approvals_needed, 1):
                self._add_rule_judgment_item(doc, i, item, 'medium')

        # 合规
        compliant = rule_judgments.get('compliant', [])
        if compliant:
            sub_heading = doc.add_heading('5.3 合规条款', level=2)
            for i, item in enumerate(compliant, 1):
                self._add_rule_judgment_item(doc, i, item, 'low')

    def _add_rule_judgment_item(self, doc: Document, index: int, item: Dict[str, Any], risk_level: str) -> None:
        """添加单个规则判定项"""
        p = doc.add_paragraph()
        p.add_run(f'{index}. ').bold = True
        p.add_run(item.get('clause_type', '条款类型'))

        if item.get('judgment'):
            p = doc.add_paragraph()
            p.add_run('判定结果：').bold = True
            run = p.add_run(item.get('judgment'))
            run.font.color.rgb = self.RISK_COLORS.get(risk_level, self.RISK_COLORS['medium'])

        if item.get('suggestion'):
            p = doc.add_paragraph()
            p.add_run('建议：').bold = True
            p.add_run(item.get('suggestion'))

        doc.add_paragraph()

    def _add_policy_references(self, doc: Document, review_result: Dict[str, Any]) -> None:
        """添加政策参考部分"""
        heading = doc.add_heading('六、适用政策参考', level=1)

        policy_refs = review_result.get('policy_references', [])
        if not policy_refs:
            doc.add_paragraph('无')
            return

        for i, policy in enumerate(policy_refs, 1):
            p = doc.add_paragraph()
            p.add_run(f'{i}. ').bold = True
            p.add_run(policy.get('policy_name', '政策'))

            if policy.get('section'):
                p = doc.add_paragraph()
                p.style = 'List Bullet'
                p.add_run(f"章节：{policy.get('section')}")

            if policy.get('content'):
                p = doc.add_paragraph()
                p.style = 'List Bullet'
                p.add_run(policy.get('content'))

        doc.add_paragraph()

    def _get_risk_level_label(self, level: str) -> str:
        """获取风险等级标签"""
        labels = {
            'high': '高风险',
            'medium': '中风险',
            'low': '低风险'
        }
        return labels.get(level, '未知')


# 全局实例
review_report_exporter = ReviewReportExporter()
