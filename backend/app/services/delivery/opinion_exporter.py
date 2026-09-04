"""Compact HTML and Word opinion exporters."""
import html
import io
from datetime import datetime
from typing import Any
from docx import Document


def generate_opinion_html(opinion: dict[str, Any]) -> bytes:
    def esc(value: Any) -> str:
        return html.escape(str(value or ""))

    sections = [f"<h1>合同审查意见</h1><p>合同：{esc(opinion.get('contract_name'))}</p>"]
    sections.append(f"<p>生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}；风险等级：{esc(opinion.get('overall_risk_level'))}</p>")
    sections.append("<h2>一、风险条款</h2><ol>")
    for item in opinion.get("risk_clauses", []):
        evidence = item.get("evidence_id", "")
        sections.append(f"<li><b>{esc(item.get('title'))}</b> [{esc(item.get('risk_level'))}]<br>原文：{esc(item.get('original_text'))}<br>意见：{esc(item.get('suggestion'))}<br>证据：{esc(evidence)}</li>")
    sections.append("</ol><h2>二、缺失条款</h2><ol>")
    for item in opinion.get("missing_clauses", []):
        sections.append(f"<li><b>{esc(item.get('title'))}</b>：{esc(item.get('description'))}<br>建议：{esc(item.get('suggestion'))}</li>")
    sections.append("</ol><h2>三、修改建议</h2><ol>")
    for item in opinion.get("recommendations", []):
        sections.append(f"<li>{esc(item.get('content') if isinstance(item, dict) else item)}</li>")
    sections.append("</ol>")
    return ("<!doctype html><html><meta charset='utf-8'><body>" + "".join(sections) + "</body></html>").encode("utf-8")


def generate_opinion_docx(opinion: dict[str, Any]) -> bytes:
    doc = Document()
    doc.add_heading("合同审查意见", level=0)
    doc.add_paragraph(f"合同：{opinion.get('contract_name', '')}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_heading("一、风险条款", level=1)
    for index, item in enumerate(opinion.get("risk_clauses", []), 1):
        doc.add_heading(f"{index}. {item.get('title', '风险条款')}", level=2)
        doc.add_paragraph(f"风险等级：{item.get('risk_level', '')}")
        doc.add_paragraph(f"原文：{item.get('original_text', '')}")
        doc.add_paragraph(f"审查意见：{item.get('risk_description', '')}")
        doc.add_paragraph(f"修改建议：{item.get('suggestion', '')}")
        if item.get("evidence_id"):
            doc.add_paragraph(f"证据 ID：{item['evidence_id']}")
    doc.add_heading("二、缺失条款", level=1)
    for item in opinion.get("missing_clauses", []):
        doc.add_paragraph(f"{item.get('title', '缺失条款')}：{item.get('description', '')}")
        doc.add_paragraph(f"建议：{item.get('suggestion', '')}")
    doc.add_heading("三、修改建议", level=1)
    for item in opinion.get("recommendations", []):
        doc.add_paragraph(str(item.get("content", item) if isinstance(item, dict) else item))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
